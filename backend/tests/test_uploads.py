"""Every supported upload format goes through create -> review with the right findings;
unsupported or broken files get a friendly 4xx instead of a crash."""
import io

import docx
import fitz
import pytest
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

from app.main import app
from app.seed_data import SAMPLES, make_poster_png

c = TestClient(app, raise_server_exceptions=False)
TEXT = SAMPLES[0]["text"].replace("₹", "Rs ")


def login(role="creator"):
    t = c.post("/api/auth/demo-login", json={"role": role}).json()
    return {"Authorization": "Bearer " + t["token"]}


def text_png(msg="Our chemical-free balm is guaranteed!") -> bytes:
    img = Image.new("RGB", (1200, 300), "white")
    ImageDraw.Draw(img).text((40, 110), msg, font=ImageFont.load_default(size=48), fill="black")
    b = io.BytesIO()
    img.save(b, "PNG")
    return b.getvalue()


def make_pdf() -> bytes:
    d = fitz.open()
    d.new_page().insert_textbox(fitz.Rect(50, 50, 550, 400), TEXT, fontsize=12)
    return d.tobytes()


def make_scanned_pdf() -> bytes:
    d = fitz.open()
    p = d.new_page(width=1200, height=300)
    p.insert_image(p.rect, stream=text_png())
    return d.tobytes()


def make_docx(with_picture=True, alt="") -> bytes:
    doc = docx.Document()
    doc.add_paragraph(TEXT)
    if with_picture:
        doc.add_picture(io.BytesIO(text_png("logo")))
        if alt:
            doc.inline_shapes[0]._inline.docPr.set("descr", alt)
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


def as_format(png: bytes, fmt: str) -> bytes:
    b = io.BytesIO()
    Image.open(io.BytesIO(png)).convert("RGB").save(b, fmt)
    return b.getvalue()


def upload(name, data, ctype, h=None, **form):
    h = h or login()
    r = c.post("/api/content", headers=h, data={"title": name, "channel": "instagram_post", **form},
               files={"file": (name, data, ctype)})
    return r, h


def review(cid, h):
    c.post(f"/api/content/{cid}/review", headers=h)
    return c.get(f"/api/content/{cid}", headers=h).json()


def rules(d):
    return {f["rule_id"] for f in d["findings"]}


TEXT_RULES = {"RULE-ABS-001", "RULE-MED-001", "RULE-TERM-001", "RULE-SUP-001"}


@pytest.mark.parametrize("name,data,ctype", [
    ("Launch post (final).pdf", make_pdf(), "application/pdf"),
    ("brief.docx", make_docx(False), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ("post.txt", TEXT.encode(), "text/plain"),
    ("post.md", TEXT.encode("utf-16"), "text/markdown"),
])
def test_text_documents_are_extracted_and_reviewed(name, data, ctype):
    r, h = upload(name, data, ctype)
    assert r.status_code == 200, r.text
    d = review(r.json()["id"], h)
    assert TEXT_RULES <= rules(d)
    assert not d["has_image"]
    for f in d["findings"]:
        if f["span"]:
            assert d["versions"][0]["raw_text"][f["span"]["start"]:f["span"]["end"]] == f["span_text"]


def test_scanned_pdf_is_read_with_ocr():
    r, h = upload("scan.pdf", make_scanned_pdf(), "application/pdf")
    assert r.status_code == 200
    d = review(r.json()["id"], h)
    assert "guaranteed" in d["versions"][0]["raw_text"].lower()
    assert d["versions"][0]["note"].startswith("Scanned PDF")
    assert {"RULE-ABS-001", "RULE-TERM-001"} <= rules(d)


def test_docx_picture_without_alt_text_is_flagged_and_with_alt_is_not():
    r, h = upload("brief.docx", make_docx(True), "application/octet-stream")
    assert "RULE-A11Y-ALT-001" in rules(review(r.json()["id"], h))
    r, h = upload("brief2.docx", make_docx(True, alt="GlowLeaf logo"), "application/octet-stream")
    assert "RULE-A11Y-ALT-001" not in rules(review(r.json()["id"], h))


def test_html_img_without_alt_is_flagged():
    r, h = upload("page.html", b'<p>Our toxin-free serum</p><img src="a.png">', "text/html")
    assert {"RULE-A11Y-ALT-001", "RULE-TERM-001"} <= rules(review(r.json()["id"], h))


@pytest.mark.parametrize("name,fmt,ctype", [
    ("poster.png", None, "image/png"),
    ("poster.JPG", "JPEG", "image/jpeg"),
    ("poster.webp", "WEBP", "image/webp"),
    ("poster.gif", "GIF", "image/gif"),
    ("poster.png", None, "application/octet-stream"),
])
def test_images_get_contrast_and_alt_checks(name, fmt, ctype):
    data = make_poster_png() if fmt is None else as_format(make_poster_png(), fmt)
    r, h = upload(name, data, ctype)
    assert r.status_code == 200, r.text
    d = review(r.json()["id"], h)
    assert d["has_image"] and d["versions"][0]["file_type"].startswith("image/")
    assert {"RULE-VIS-CONTRAST-001", "RULE-A11Y-ALT-001"} <= rules(d)
    img = c.get(f"/api/content/{d['content']['id']}/image", headers=h)
    assert img.status_code == 200 and img.headers["content-type"].startswith("image/")


def test_adding_alt_text_in_a_new_version_clears_the_finding():
    r, h = upload("poster.png", make_poster_png(), "image/png")
    cid = r.json()["id"]
    review(cid, h)
    d = c.post(f"/api/content/{cid}/versions", headers=h, json={"alt_text": "Summer sale poster, 50% off lip care"}).json()
    assert d["versions"][-1]["alt_text"].startswith("Summer sale")
    assert "RULE-A11Y-ALT-001" not in rules(d) and "RULE-VIS-CONTRAST-001" in rules(d)


@pytest.mark.parametrize("name,data,ctype,status", [
    ("bad.pdf", b"not a pdf", "application/pdf", 400),
    ("bad.docx", b"PK\x03\x04junk", "application/octet-stream", 400),
    ("bad.png", b"not an image", "image/png", 415),
    ("deck.pptx", b"PK\x03\x04junk", "application/vnd.ms-powerpoint", 415),
    ("old.doc", b"\xd0\xcf\x11\xe0", "application/msword", 415),
    ("photo.heic", b"\x00\x00\x00\x18ftypheic", "image/heic", 415),
    ("archive.zip", b"PK\x03\x04", "application/zip", 415),
    ("empty.txt", b"", "text/plain", 400),
])
def test_bad_or_unsupported_files_get_a_friendly_error(name, data, ctype, status):
    r, _ = upload(name, data, ctype)
    assert r.status_code == status
    assert "went wrong" not in r.json()["detail"] and len(r.json()["detail"]) > 20


def test_oversized_file_rejected():
    r, _ = upload("big.txt", b"a" * (10 * 1024 * 1024 + 1), "text/plain")
    assert r.status_code == 413


def test_storage_paths_are_unique_and_safe():
    h = login()
    a, _ = upload("my post #1 (final).txt", b"Hello there you", "text/plain", h)
    b, _ = upload("my post #1 (final).txt", b"Different text you", "text/plain", h)
    pa = c.get(f"/api/content/{a.json()['id']}", headers=h).json()["versions"][0]["file_path"]
    pb = c.get(f"/api/content/{b.json()['id']}", headers=h).json()["versions"][0]["file_path"]
    assert pa != pb and pa.endswith(".txt") and all(ch not in pa.split("/")[-1] for ch in " #()")


def test_admin_can_add_sources_from_files():
    h = login("admin")
    for name, data, ctype in [("facts.pdf", make_pdf(), "application/pdf"), ("terms.docx", make_docx(False), "application/octet-stream"),
                              ("scan.png", text_png(), "image/png"), ("notes.txt", b"Approved claim text", "text/plain")]:
        r = c.post("/api/sources", headers=h, data={"kind": "guideline", "title": name}, files={"file": (name, data, ctype)})
        assert r.status_code == 200 and r.json()["chunk_count"] >= 1, (name, r.text)
    r = c.post("/api/sources", headers=h, data={"kind": "guideline", "title": "x"}, files={"file": ("x.pptx", b"PK", "application/octet-stream")})
    assert r.status_code == 415
